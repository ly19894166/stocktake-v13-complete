"""卓不凡专项核查；只读数量金额，人工结论使用已有审计字段。"""
import io
from decimal import Decimal as D
from collections import defaultdict
from core import num,current_total,newest_diff,current_amount,valid_count,patch_items
from supplier_policy import SupplierDirectory,effective_vendors,CANONICAL

HEADERS=('序号','品类','商品名称','商品条码','商品编码','供应商','系统数','初盘数','复盘实数','当前总数','最新差异','差异成本','疑似A换B对应商品','人工备注','核查结论','凭证编号')

def audit_model(items,vendors,meta):
    directory=SupplierDirectory(vendors,meta);classes=directory.classify_all(items)
    groups=defaultdict(lambda:{1:[],-1:[]});byid={x['记录ID']:x for x in items}
    for x in items:
        c=classes[x['记录ID']];d=newest_diff(x)
        if c['status']=='zbf' and d not in (None,0):groups[c['code']][1 if d>0 else -1].append(x)
    peers={}
    for x in items:
        c=classes[x['记录ID']];d=newest_diff(x)
        peers[x['记录ID']]=[y for y in groups[c['code']][-1 if d>0 else 1] if y['商品条码']!=x['商品条码']] if c['status']=='zbf' and d not in (None,0) else []
    def verified(x):return bool(str(x.get('核实结论') or '').strip() and str(x.get('相关单号') or '').strip())
    def proof(a,b):
        key='|'.join(sorted((str(a['序号']),str(b['序号']))))
        return meta.get('pair_proofs',{}).get(key,{})
    confirmed=[x for x in items if classes[x['记录ID']]['status']=='zbf']
    def amounts(rows):
        values=[current_amount(x) for x in rows];known=[v for v in values if v is not None]
        return {'盘盈':sum((v for v in known if v>0),D(0)),'盘亏':sum((v for v in known if v<0),D(0)),'净差异':sum(known,D(0)),'缺失金额':len(values)-len(known)}
    rows=[]
    for x in items:
        c=classes[x['记录ID']]
        if c['status'] not in ('zbf','pending'):continue
        proofs=[proof(x,y) for y in peers[x['记录ID']]]
        conclusion='；'.join(v for v in [x.get('核实结论',''),meta.get('zbf_notes',{}).get(x['记录ID'],''),*[p.get('说明','') for p in proofs]] if v)
        documents='；'.join(dict.fromkeys(v for v in [x.get('相关单号',''),*[p.get('单号','') for p in proofs]] if v))
        rows.append(dict(zip(HEADERS,[x['序号'],x.get('大类') or '未分类',x['商品名称'],x['商品条码'],x['商品编码'],c['name'],num(x['系统数']),num(x['初盘数']),num(x.get('复盘实数')),num(current_total(x)),num(newest_diff(x)),num(current_amount(x)),
            '；'.join(y['商品名称']+' / '+y['商品条码'] for y in peers[x['记录ID']]),x.get('备注',''),conclusion,documents]))|{'_id':x['记录ID'],'_status':c['status'],'_verified':verified(x)})
    edges={tuple(sorted((a,y['记录ID']))) for a,ys in peers.items() for y in ys}
    unresolved=sum(not (verified(byid[a]) and verified(byid[b]) or proof(byid[a],byid[b]).get('说明') and proof(byid[a],byid[b]).get('单号')) for a,b in edges)
    stats={'商品总条数':len(confirmed),'已复盘':sum(valid_count(x.get('复盘实数')) for x in confirmed),'未复盘':sum(not valid_count(x.get('复盘实数')) for x in confirmed),**amounts(confirmed),'疑似A换B对数':len(edges),'未核实对数':unresolved,'供货商待确认':sum(c['status']=='pending' for c in classes.values())}
    comparisons={name:amounts(subset) for name,subset in [('含卓不凡全口径',items),('已确认卓不凡',confirmed),('非卓不凡内部分析（不含待确认）',[x for x in items if classes[x['记录ID']]['status']=='other']),('归属待确认',[x for x in items if classes[x['记录ID']]['status']=='pending'])]}
    return {'rows':rows,'stats':stats,'comparisons':comparisons,'edges':edges,'classes':classes,'byid':byid}

def filtered_rows(model,status='全部',category='全部',include_pending=False):
    rows=model['rows']
    return [r for r in rows if (r['_status']=='zbf' or include_pending or status=='归属待确认') and (category=='全部' or r['品类']==category) and
        (status=='全部' or status=='归属待确认' and r['_status']=='pending' or status=='未复盘' and r['复盘实数'] is None or status=='已复盘' and r['复盘实数'] is not None or status=='仍有差异' and r['最新差异'] not in (None,0) or status=='疑似A换B' and r['疑似A换B对应商品'] or status=='待补凭证' and r['疑似A换B对应商品'] and not r['_verified'])]

def audit_excel(model,rows):
    import openpyxl
    w=openpyxl.Workbook();s=w.active;s.title='卓不凡专项审计';s.append(HEADERS)
    for row in rows:
        s.append([row[k] for k in HEADERS])
        for col in (4,5):s.cell(s.max_row,col).data_type='s';s.cell(s.max_row,col).number_format='@'
        # 外部备注/名称按文本写入，不能成为Excel公式。
        for cell in s[s.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    s.freeze_panes='F2';s.auto_filter.ref=s.dimensions
    for col,width in [('A',8),('B',18),('C',30),('D',21),('E',18),('F',18),('M',65),('N',35),('O',35),('P',24)]:s.column_dimensions[col].width=width
    stats=w.create_sheet('汇总与口径');stats.append(['项目','数值'])
    for k,v in model['stats'].items():stats.append([k,float(v) if isinstance(v,D) else v])
    stats.append(['说明','专项审计不抵销商品，不修改公司三列数量或整店2‰正式口径；未核实金额不等于已确认损耗。'])
    for name,values in model['comparisons'].items():stats.append([name,'；'.join(k+'='+str(v) for k,v in values.items())])
    pairs=w.create_sheet('A换B线索');pairs.append(['商品A条码','商品A名称','商品B条码','商品B名称','提示'])
    shown={r['_id'] for r in rows}
    for a,b in sorted(model['edges']):
        if a not in shown and b not in shown:continue
        x,y=model['byid'][a],model['byid'][b];pairs.append([x['商品条码'],x['商品名称'],y['商品条码'],y['商品名称'],'疑似待核实，不抵销、不删除'])
        for cell in pairs[pairs.max_row]:cell.data_type='s'
    fp=io.BytesIO();w.save(fp);return fp.getvalue()

def audit_panel(db,batch,items,state,rev):
    import streamlit as st
    outcome=st.session_state.get('zbf_result_'+batch)
    if outcome:(st.success if outcome[0] else st.error)(outcome[1])
    pane=st.expander('卓不凡专项审计',on_change='rerun')
    if not pane.open:return
    with pane:
        model=audit_model(items,effective_vendors(state),state['meta']);stats=model['stats']
        st.write(f"卓不凡：{stats['商品总条数']}条 · 已复盘{stats['已复盘']} / 未复盘{stats['未复盘']} · 供货商待确认{stats['供货商待确认']}条")
        st.write(f"盘盈{stats['盘盈']:.2f}元 · 盘亏{stats['盘亏']:.2f}元 · 净差异{stats['净差异']:.2f}元 · 疑似A换B {stats['疑似A换B对数']}对 / 未核实{stats['未核实对数']}对")
        if stats['缺失金额']:st.warning(f"{stats['缺失金额']}条成本无法计算，汇总仅含已知金额。")
        a,b=st.columns(2)
        status=a.selectbox('专项审计状态',['全部','未复盘','已复盘','仍有差异','疑似A换B','待补凭证','归属待确认'],key='zbf_status_'+batch)
        cat=b.selectbox('专项审计品类',['全部']+sorted({r['品类'] for r in model['rows']}),key='zbf_cat_'+batch)
        pending=st.checkbox('专项表同时显示供货商待确认商品',key='zbf_pending_'+batch)
        rows=filtered_rows(model,status,cat,pending)
        st.dataframe([{k:r[k] for k in HEADERS} for r in rows],hide_index=True,alt='卓不凡独立专项审计，不是公司上传表')
        st.download_button('下载卓不凡专项审计Excel',audit_excel(model,rows),'卓不凡专项审计_非公司上传.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        confirmed=[r for r in model['rows'] if r['_status']=='zbf']
        if confirmed:
            selected=st.selectbox('保留人工核查结论及凭证的商品',[r['_id'] for r in confirmed],format_func=lambda k:model['byid'][k]['商品名称']+' / '+model['byid'][k]['商品条码'])
            x=model['byid'][selected]
            token=batch+'_'+selected
            def save_proof():
                from storage import Store
                live=Store(db.root)
                try:
                    fresh,_,latest,version=live.load(batch)
                    if version!=rev:raise ValueError('另一窗口已更新，请刷新后核对专项凭证，未覆盖原记录。')
                    latest['meta'].setdefault('zbf_notes',{})[selected]=st.session_state['zbf_note_'+token]
                    updated=patch_items(fresh,{selected:{'核实结论':st.session_state['zbf_conclusion_'+token],'相关单号':st.session_state['zbf_doc_'+token]}})
                    live.save(batch,updated,latest,version,'本机店长','卓不凡专项说明与凭证')
                    st.session_state['zbf_result_'+batch]=(True,'专项说明与凭证已保存')
                except ValueError as exc:st.session_state['zbf_result_'+batch]=(False,str(exc))
                finally:live.close()
            with st.form('zbf_proof_'+batch):
                conclusions=['','已排除','无需调整','确认并附单据']
                st.selectbox('核查结论（沿用已有结论）',conclusions,index=conclusions.index(x.get('核实结论','')),key='zbf_conclusion_'+token)
                st.text_area('专项核查说明（人工）',value=state['meta'].get('zbf_notes',{}).get(selected,''),key='zbf_note_'+token)
                st.text_input('凭证编号',value=x.get('相关单号',''),key='zbf_doc_'+token)
                st.form_submit_button('保存专项说明与凭证',on_click=save_proof)
        with st.expander('内部分析口径对照'):
            st.caption('非卓不凡口径仅供内部分析，不代表公司最终损耗结论。公司导出和顶部2‰始终包含整店商品。')
            st.dataframe([{'口径':n,**{k:num(v) if isinstance(v,D) else v for k,v in values.items()}} for n,values in model['comparisons'].items()],hide_index=True,alt='全口径与仅内部分析口径')
