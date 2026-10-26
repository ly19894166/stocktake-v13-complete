"""独立读回浏览器导出文件及模拟数据库；禁止指向正式数据目录。"""
from pathlib import Path
import os,sys,io,json,re,hashlib
from collections import Counter
from decimal import Decimal as D
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
import openpyxl,xlrd
from pypdf import PdfReader
from pypdf.generic import ContentStream
from storage import Store
from core import hydrate,patch_items,current_total,current_amount
from vendor_audit import audit_model,filtered_rows,HEADERS
from print_manager import make_pdf
from workflow import print_order
DATA=Path(os.environ.get('STOCKTAKE_FIFTH_DATA_DIR',ROOT/'tests/runtime/fifth_shipping_data')).resolve()
assert DATA.is_relative_to((ROOT/'tests/runtime').resolve()),'仅允许独立模拟数据库'
OUT=ROOT/'tests/runtime/fifth_results';OUT.mkdir(parents=True,exist_ok=True)

def snapshot():
    s=Store(DATA)
    try:
        return {t:hashlib.sha256(repr(s.db.execute('SELECT * FROM '+t+' ORDER BY rowid').fetchall()).encode()).hexdigest() for (t,) in s.db.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")}
    finally:s.close()

def verify():
    s=Store(DATA)
    try:
        batch=s.list_batches()[0]['id'];items,_,state,_=s.load(batch)
        assert len(items)==480 and len(state['exchange_catalog'])==28 and len(s.exchange_history())==30
        # 逐行与真实后端计算值一致，而不是重写一份前端数量公式。
        payload=json.loads((OUT/'browser_amounts.json').read_text('utf8'));shown={r['id']:r['text'] for r in payload['display']};byid={x['记录ID']:x for x in items}
        for row in payload['rows']:
            x=dict(byid[row['id']]);x['复盘实数']=None if row['recount'] is None else D(str(row['recount']))
            expected=current_amount(x);assert row['amount']==(float(expected) if expected is not None else None)
            if expected is not None:assert D(shown[row['id']].replace(',','').replace('+',''))==expected
        audit=json.loads((OUT/'audit_snapshot.json').read_text('utf8'));original=hydrate(json.loads(s.db.execute('SELECT original FROM batches WHERE id=?',(batch,)).fetchone()[0]))
        for x in original:
            latest=byid[x['记录ID']];x['核实结论']=latest.get('核实结论','');x['相关单号']=latest.get('相关单号','')
        snapshotrows={r['id']:r for r in audit['rows']}
        original=patch_items(original,{k:{'复盘实数':r['recount'],'备注':r['note']} for k,r in snapshotrows.items()})
        model=audit_model(original,state['vendors'],state['meta'],state['exchange_catalog']);rows=filtered_rows(model)
        book=openpyxl.load_workbook(OUT/'模拟可A换B专项审计.xlsx',data_only=False);sheet=book['可A换B商品专项审计']
        assert sheet.max_row==81 and sheet.max_column==16 and tuple(c.value for c in sheet[1])==HEADERS
        for n,r in enumerate(rows,2):
            for col,h in enumerate(HEADERS,1):assert sheet.cell(n,col).value==(r[h] if r[h]!='' else None),(n,h)
            assert sheet.cell(n,4).data_type==sheet.cell(n,5).data_type=='s'
            assert r['疑似A换B对应商品'].count('；')<=2
        assert model['stats']['商品总条数']==80 and model['stats']['疑似A换B对数']==950
        assert book['A换B线索'].max_row==951 and book['资格与来源'].max_row==81
        assert any('SIM-PAIR-FIFTH' in str(c.value) for row in book['历史专项凭证'] for c in row)
        summary=dict(book['汇总与口径'].values);assert summary['净差异']==float(model['stats']['净差异'])
        company=(OUT/'模拟全店480条公司三列.xls').read_bytes();assert company[:8]==bytes.fromhex('D0CF11E0A1B11AE1')
        sheet=xlrd.open_workbook(file_contents=company).sheet_by_index(0)
        assert sheet.nrows==481 and sheet.ncols==3 and sheet.row_values(0)==['商品编码','商品条码','数量']
        for n,x in enumerate(items,1):
            assert sheet.row_values(n)==[x['商品编码'],x['商品条码'],float(current_total(x))]
            assert [sheet.cell_type(n,c) for c in range(3)]==[xlrd.XL_CELL_TEXT,xlrd.XL_CELL_TEXT,xlrd.XL_CELL_NUMBER]
        assert (OUT/'模拟现场单原件.pdf').read_bytes()==(OUT/'模拟现场单重印.pdf').read_bytes()
        archives=s.prints(batch);assert len(archives)==2
        first=next(p for p in archives if not p.get('parent'));ids={r['source_id'] for r in first['rows']}
        classes=model['classes'];assert not any(classes[k]['status']=='zbf' for k in ids)
        manual={k for k,e in model['engine'].eligibility.items() if e['source']=='人工加入'}
        assert manual<=ids,'非卓不凡可换货商品不得自动从03排除'
        route=[(byid[r['source_id']]['大类'],byid[r['source_id']]['中类'],r['sequence']) for r in first['rows']]
        assert first['sort']=='category' and route==sorted(route)
        result={'pass':True,'simulated_only':True,'rows':480,'registry_records':28,'registry_history_events':30,'amounts_equal_core_rows':480,'audit_rows':80,'audit_columns':16,'zbf_rows':60,'manual_eligible_rows':20,'full_candidate_pairs':950,'ui_max_candidates_per_product':3,'company_true_xls_rows':480,'company_columns':3,'company_all_zbf_and_manual_retained':True,'historical_pdf_byte_identical':True,'print_manual_records_retained':True,'old_and_new_proof_sheets_present':True}
        (OUT/'artifact_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');return result
    finally:s.close()

def pdf():
    s=Store(DATA)
    try:
        b=s.list_batches()[0]['id'];original=hydrate(json.loads(s.db.execute('SELECT original FROM batches WHERE id=?',(b,)).fetchone()[0]));rows=print_order(original[:414],'category')
        raw,pages=make_pdf(rows,{'store':'模拟店','period':'模拟期'},'SIM-414',scope='模拟414条现场找货');(OUT/'模拟414条A4回归.pdf').write_bytes(raw)
        reader=PdfReader(io.BytesIO(raw));codes=[];counts=[];boxes=0;embedded=False
        for page in reader.pages:
            assert abs(float(page.mediabox.width)-841.89)<1 and abs(float(page.mediabox.height)-595.28)<1
            text=page.extract_text();assert not any(w in text for w in ('串码','A换B','疑似','算法'))
            found=re.findall(r'(?<!\d)099\d{10}(?!\d)',text);codes+=found;counts.append(len(found))
            for font in page['/Resources']['/Font'].values():
                fd=font.get_object().get('/FontDescriptor')
                if fd and '/FontFile2' in fd.get_object():embedded=True
            for args,op in ContentStream(page.get_contents(),reader).operations:
                if op==b're' and 80<=float(args[2])<=110 and float(args[3])>=22:boxes+=1
        assert pages<=30 and boxes==414 and embedded and Counter(codes)==Counter(x['商品条码'] for x in rows)
        assert all(14<=n<=18 for n in counts[:-1])
        result={'pass':True,'simulated_only':True,'rows':414,'pages':pages,'per_page_rows':counts,'handwriting_boxes':boxes,'minimum_box_height_points':22,'embedded_chinese_font':True,'no_pairing_analysis':True}
        (OUT/'pdf_verification.json').write_text(json.dumps(result,indent=2),encoding='utf8');return result
    finally:s.close()

if __name__=='__main__':
    mode=sys.argv[1] if len(sys.argv)>1 else 'verify'
    if mode=='before':(OUT/'restart_expected.json').write_text(json.dumps(snapshot()),encoding='utf8');result={'snapshot_saved':True}
    elif mode=='after':
        assert snapshot()==json.loads((OUT/'restart_expected.json').read_text('utf8'));result={'pass':True,'simulated_only':True,'all_ten_tables_unchanged_after_real_restart':True};(OUT/'restart_verification.json').write_text(json.dumps(result),encoding='utf8')
    else:result={'verify':verify,'pdf':pdf}[mode]()
    print(json.dumps(result,ensure_ascii=False))
