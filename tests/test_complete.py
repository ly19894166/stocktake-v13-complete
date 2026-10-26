import copy,io,json,time,sqlite3,subprocess,sys,hashlib
from pathlib import Path
from decimal import Decimal as D
import pytest
import openpyxl
from core import (patch_items,newest_diff,current_total,current_amount,summary,
    export_blockers,default_vendors,create_company_xls,validate_all)
from intelligence import CandidateEngine,row_models,unit_hints,priority
from imports_exports import import_master,supplier_import,generic_rows,cross_candidates,transactions_for,export_workbook
from imports_exports import identify_batch
from storage import Store

def rows_from(items,n=2):
    rows=[]
    for i in range(n):
        x=copy.deepcopy(items[0]);x.update({'记录ID':f'test:{i}','序号':i+1,'商品编码':f'c{i}',
            '商品条码':f'069000000{i:05d}','商品名称':f'独立商品{i}',
            '部类':'水饮','大类':'饮料','中类':'茶饮','小类':'乌龙',
            '成本价':D('3.50'),'零售价':D('5.00'),'系统数':D(20 if i==0 else 10),
            '初盘数':D(10 if i==0 else 20),'原始调整':D(0),'复盘实数':None,
            '备注':'人工原因不可删除','品牌':'','包装系数':None,'导入问题':[]})
        rows.append(x)
    return rows

@pytest.mark.parametrize('variant,expected',[
    ('same',1),('different_middle',1),('only_major',1),('only_middle',1),
    ('no_category',0),('different_cost',0),('same_direction',0),('same_barcode',0),
    ('no_retail',1),('different_retail',1),('small_difference',1),('round_cents',1)])
def test_latest_candidate_conditions(items,variant,expected):
    rows=rows_from(items);a,b=rows
    if variant=='different_middle':b.update(大类='其他大类',中类='其他中类')
    elif variant=='only_major':b.update(部类='其他',中类='其他')
    elif variant=='only_middle':b.update(部类='其他',大类='其他')
    elif variant=='no_category':b.update(部类='其他',大类='其他',中类='其他')
    elif variant=='different_cost':b['成本价']=D('3.51')
    elif variant=='same_direction':b['初盘数']=D(0)
    elif variant=='same_barcode':b['商品条码']=a['商品条码']
    elif variant=='no_retail':b['零售价']=None
    elif variant=='different_retail':b['零售价']=D('99')
    elif variant=='small_difference':b['初盘数']=D('10.1')
    elif variant=='round_cents':b['成本价']=D('3.504')
    e=CandidateEngine(rows,[])
    assert e.count(a['记录ID'])==expected
    if variant=='different_cost':assert CandidateEngine(rows,[],{'cost_tolerance':True}).count(a['记录ID'])==1

def test_multiple_candidates_zero_only_removes_related_pairs(items):
    rows=rows_from(items,4);a,b,c,d=rows
    e=CandidateEngine(rows,[]);assert e.count(a['记录ID'])==3
    b['复盘实数']=b['系统数']
    e=CandidateEngine(rows,[])
    assert e.count(a['记录ID'])==2 and e.count(b['记录ID'])==0
    assert c['记录ID'] in e.neighbors(a['记录ID'])
    assert all(x['备注']=='人工原因不可删除' for x in rows)
    b['复盘实数']=b['初盘数']
    assert CandidateEngine(rows,[]).count(a['记录ID'])==3

def test_supplier_swap_main_choice_and_ranking(items):
    rows=rows_from(items,4);a,b,c,d=rows
    b.update(部类='其他',大类='其他',中类='其他',成本价=D(99))
    vendors=[{'商品条码':x['商品条码'],'供货商名称':'模拟供货商甲'} for x in rows]
    e=CandidateEngine(rows,vendors)
    assert b['记录ID'] not in e.neighbors(a['记录ID']) and not e.neighbors(a['记录ID'],'swap')
    assert e.peers(a['记录ID'])[0]['mix'] is True
    vendors.append({'商品条码':a['商品条码'],'供货商名称':'另一家'})
    assert not CandidateEngine(rows,vendors).neighbors(a['记录ID'],'swap')
    assert CandidateEngine(rows,vendors,main={a['商品条码']:'模拟供货商甲'}).count(a['记录ID'],'swap')==0 # 同供应商不再自动授权

def test_twenty_thousand_rows_indexed_count_and_validation(items):
    # 高度同类同成本也不能为UI计数物化一亿条边。
    seed=rows_from(items,2);rows=[]
    for i in range(20000):
        x=copy.copy(seed[i%2]);x.update(记录ID=f'perf:{i}',序号=i+1,商品编码=str(i),商品条码=f'690{i:010d}')
        rows.append(x)
    start=time.perf_counter();e=CandidateEngine(rows,[])
    assert e.count(rows[0]['记录ID'])==10000
    state={'meta':{'parameters':{}},'vendors':[]}
    models=row_models(rows,state,e,enrich={x['记录ID'] for x in rows[:20]})
    assert len(models)==20000 and models[-1]['mix_count']==10000
    blockers=export_blockers(rows,{'both_files':True},[],engine=e)
    assert blockers
    elapsed=time.perf_counter()-start
    assert elapsed<20,f'2万条计数/筛选模型/审核耗时{elapsed:.2f}s'


def excel(headers,rows):
    w=openpyxl.Workbook();s=w.active;s.append(headers)
    for r in rows:s.append(r)
    out=io.BytesIO();w.save(out);return out.getvalue()

def test_supplier_bulk_mapping_and_all_rows_preserved():
    blob=excel(['国际条码','厂商','商品名'],[['0012345678901','模拟供货商甲','商品A'],['0012345678901','另供货商','商品A']])
    rows=supplier_import(blob,'供应商.xlsx')
    assert len(rows)==2 and all(r['商品条码']=='0012345678901' for r in rows)
    blob=excel(['自定条码','自定供货方'],[['0012345678902','某供应商']])
    assert supplier_import(blob,'自定义.xlsx',{'商品条码':'自定条码','供货商名称':'自定供货方'})[0]['商品条码']=='0012345678902'

def test_optional_cross_store_and_transactions_do_not_change_stock(items):
    rows=rows_from(items);before=copy.deepcopy(rows)
    a=rows[0];bc=a['商品条码']
    cross=generic_rows(excel(['商品条码','差异数量','日期','单号'],[[bc,10,'2026-10-22','T01']]),'跨店.xlsx',('商品条码','差异数量'))
    sales=generic_rows(excel(['商品条码','销售数量','销售日期'],[[bc,2,'2026-10-23']]),'204.xlsx')
    evidence=[{'类型':'跨店差异','门店':'模拟门店乙','文件':'跨店.xlsx','rows':cross},
              {'类型':'销售明细204','门店':'','文件':'204.xlsx','rows':sales}]
    assert cross_candidates(rows,evidence)[a['记录ID']][0]['门店']=='模拟门店乙'
    assert transactions_for(bc,evidence,'2026-10-23','2026-10-23')[0]['数量']=='2'
    assert not transactions_for(bc,evidence,'2026-10-24','2026-10-25')
    assert rows==before

def test_unit_issues_filters_and_empty_no_invented_factor(items):
    rows=rows_from(items);x=rows[0];x.update(系统数=D(12),初盘数=D(0),包装系数=D(6))
    assert any('包装系数6' in h for h in unit_hints(x))
    x.update(系统数=D(-1),成本价=None,商品编码='',原始调整=D('.4'))
    models=row_models(rows,{'vendors':[],'meta':{}},enrich=[])
    assert {'负库存','缺成本','缺编码','异常换算'}<=set(models[0]['flags'])
    x.update(系统数=D(10),初盘数=D(20),包装系数=None,销售规格='')
    assert any('仅作猜测' in h for h in unit_hints(x))
    assert '建议排序' in priority(x)[1]

def test_atomic_undo_idempotence_and_full_reopen(store,batch):
    items,_,state,rev=store.load(batch);x=items[0]
    updated=patch_items(items,{x['记录ID']:{'复盘实数':0,'备注':'人工备注'}})
    rev=store.save(batch,updated,state,rev,'店长','批量测试','same-request')
    assert store.save(batch,updated,state,rev-1,'店长','批量测试','same-request')==rev
    assert len(store.item_history(batch))==1
    assert store.item_history(batch)[0]['修改来源']=='批量测试'
    store.undo_last(batch,rev)
    loaded=store.load(batch)[0][0]
    assert loaded['复盘实数'] is None and loaded['备注']==x['备注']
    assert store.item_history(batch)[-1]['已撤销']

def test_evidence_blobs_and_unknown_store_no_required_form(store,sample):
    batch=store.create(*sample[:3],sources=sample[3]);items,_,state,rev=store.load(batch)
    assert state['meta']['store']=='门店未识别' and state['meta']['period']=='期间未识别'
    assert '+08:00' in state['meta']['first_import_at']
    assert store.create(*sample[:3],sources=sample[3])==batch
    blob=excel(['商品条码','数量'],[['000123',2]])
    rows=generic_rows(blob,'销售.xlsx')
    ident=store.add_evidence(batch,'销售明细204','','销售.xlsx',blob,rows)
    assert store.add_evidence(batch,'销售明细204','','销售.xlsx',blob,rows)==ident
    assert store.evidence(batch)[0]['rows'][0]['商品条码']=='000123'

def test_v13_full_export_seven_sheets_warning_and_text_ids(items):
    rows=rows_from(items,3);state={'vendors':default_vendors(),'meta':{'store':'门店未识别','period':'期间未识别'}}
    e=CandidateEngine(rows,state['vendors']);blob=export_workbook(rows,state,e)
    w=openpyxl.load_workbook(io.BytesIO(blob),data_only=False)
    assert len(w.sheetnames)==7
    assert '草稿' in w['01现场复盘']['A1'].value
    assert w['01现场复盘']['D3'].data_type=='s'
    assert w['01现场复盘']['D3'].value==rows[0]['商品条码']
    assert '疑似' in w['01现场复盘']['M3'].value
    assert w['02自动清单与系统复制']['C2'].value=='数量'

def test_reliable_metadata_only_and_multiple_supplier_barcodes():
    b=excel(['商品条码','门店名称','商品名称'],[['00123','模拟门店甲','商品1'],['00124','模拟门店甲','商品2']])
    assert identify_batch({'原件':('源.xlsx',b)})['store']=='模拟门店甲'
    b=excel(['商品条码','商品名称'],[['00123','商品1']])
    assert 'store' not in identify_batch({'原件':('南上2026Q3.xlsx',b)})
    b=excel(['商品条码','供应商','副条码1'],[['0012345678901','某供应商','0012345678902']])
    assert {r['商品条码'] for r in supplier_import(b,'多条码.xlsx')}=={'0012345678901','0012345678902'}

def test_xls_evidence_dates_and_zero_display_precision():
    import xlwt
    from datetime import datetime
    w=xlwt.Workbook();s=w.add_sheet('销售')
    for c,v in enumerate(['商品条码','日期','数量']):s.write(0,c,v)
    s.write(1,0,'0012345678901');s.write(1,1,datetime(2026,10,23),xlwt.easyxf(num_format_str='YYYY-MM-DD'));s.write(1,2,2)
    b=io.BytesIO();w.save(b)
    assert generic_rows(b.getvalue(),'日期.xls')[0]['时间'].startswith('2026-10-23')
    w=xlwt.Workbook();s=w.add_sheet('供应商');s.write(0,0,'商品条码');s.write(0,1,'供应商')
    s.write(1,0,123,xlwt.easyxf(num_format_str='0000000000000'));s.write(1,1,'某供应商')
    b=io.BytesIO();w.save(b)
    with pytest.raises(ValueError):supplier_import(b.getvalue(),'格式前导零.xls')
