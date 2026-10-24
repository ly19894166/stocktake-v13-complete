"""独立手算模拟输入；此测试不声称是公司真实文件/WPS缓存回归。"""
import io
from decimal import Decimal as D, ROUND_HALF_UP
import pytest
import openpyxl
from core import current_amount,newest_diff,current_total,initial_diff,summary,validate_all
from imports_exports import import_master
from synthetic_data import source_rows, simulated_master

def test_simulated_source_every_row_matches_independent_arithmetic(items):
    expected={x['商品条码']:x for x in source_rows()}
    assert len(items)==len(expected)==12
    assert not validate_all(items)
    total=D(0)
    for x in items:
        source=expected[x['商品条码']]
        diff=D(source['初盘数量'])-D(source['库存数量'])+D(source['子转母商品数量'])
        amount=(diff*D(str(source['成本价']))).quantize(D('.01'),rounding=ROUND_HALF_UP)
        assert x['商品编码']==source['商品'] and x['商品名称']==source['商品名称']
        assert newest_diff(x)==initial_diff(x)==diff
        assert current_total(x)==source['初盘数量']
        assert current_amount(x)==amount
        total+=amount
    assert summary(items,None)['净差异']==total

@pytest.mark.parametrize('count',[7,37,53])
def test_simulated_master_dynamic_range_and_original_conversion(count):
    blob=simulated_master(count)
    rows,notes,fp,vendors=import_master(blob,'模拟V13母版.xlsx')
    w=openpyxl.load_workbook(io.BytesIO(blob));s=w['01现场复盘'];a=w['04商品档案后台']
    assert len(rows)==count and len(vendors)==1
    assert vendors[0]['供货商名称']=='模拟供货商甲'
    assert rows[0]['原始调整']==D('.4')
    for index,x in enumerate(rows,6):
        expected=D(str(s.cell(index,7).value))-D(str(s.cell(index,6).value))+D(str(a.cell(index,13).value))
        assert newest_diff(x)==expected
        assert current_amount(x)==(expected*D(str(a.cell(index,11).value))).quantize(D('.01'),rounding=ROUND_HALF_UP)
    rows[0]['复盘实数']=D(0)
    assert newest_diff(rows[0])==D('.4')-rows[0]['系统数']
    w.close()
