"""公开测试全部使用程序生成的模拟样本，无公司业务数据。"""
import copy
import io
import json
from decimal import Decimal as D
from pathlib import Path
import openpyxl
import pytest
import xlrd
import xlwt
from core import (summary,initial_diff,newest_diff,current_amount,current_total,
    validate_all,changes,candidate_counts,legacy_candidate_counts,default_vendors,supplier_map,patch_items,
    export_blockers,create_company_xls,import_pair,import_v13,read_source)

FIX=Path(__file__).parent/'fixtures'


def test_stale_source_differences_not_used_as_adjustment(items):
    affected=[x for x in items if x['源报差异数']!=initial_diff(x)]
    assert affected
    for x in affected:
        assert x['原始调整']==(x['子转母'] or 0)
        assert newest_diff(x)!=x['源报差异数']

def test_blank_zero_clear_and_precision(items):
    x=items[0];original=current_total(x)
    x['复盘实数']=D(0)
    assert current_total(x)==0
    assert summary(items,None)['已复盘']==1
    x['复盘实数']=None
    assert current_total(x)==original
    x['复盘实数']=D('1.23455')
    assert current_total(x)==D('1.2346')
    x['成本价']=D('.5');x['系统数']=D(0);x['原始调整']=D(0);x['复盘实数']=D('1.01')
    assert current_amount(x)==D('.51')

def test_found_surplus_and_direction_crossing(items):
    short=next(x for x in items if initial_diff(x)<-1)
    base=initial_diff(short)
    short['复盘实数']=short['初盘数']+1
    assert changes(short)['找回数量']==1
    short['复盘实数']=short['系统数']-short['原始调整']+2
    assert changes(short)['找回数量']==-base
    assert changes(short)['盘盈增加量']==2
    surplus=next(x for x in items if initial_diff(x)>1)
    old=initial_diff(surplus)
    surplus['复盘实数']=surplus['初盘数']-1
    assert changes(surplus)['盘盈减少量']==1
    surplus['复盘实数']=surplus['初盘数']+1
    assert changes(surplus)['盘盈增加量']==1
    surplus['复盘实数']=D(0)
    assert changes(surplus)['盘盈减少量']<=old

def pair(items):
    a,b=copy.deepcopy(items[:2])
    for i,x in enumerate((a,b)):
        x['序号']=i+1;x['系统数']=D(2 if i==0 else 1)
        x['初盘数']=D(1 if i==0 else 2);x['复盘实数']=None;x['原始调整']=D(0)
        x['成本价']=D('1.11');x['零售价']=D(3)
        x.update({'部类':'测试','大类':'测试','中类':'测试'})
    vendors=[{'商品条码':x['商品条码'],'供货商名称':'测试供应商'} for x in (a,b)]
    return [a,b],vendors

def test_price_filter_and_zero_cancels_both_candidates(items):
    group,vendors=pair(items);a,b=group
    for c in candidate_counts(group,vendors).values():assert len(c['串码'])==1 and len(c['换货'])==0 # 普通供应商无换货授权
    b['零售价']=D(4)
    # 最新要求：售价只作证据与排序，不阻止疑似匹配。
    for c in candidate_counts(group,[]).values():assert c['串码']
    b['零售价']=None
    for c in candidate_counts(group,[]).values():assert c['串码']
    a['复盘实数']=D(2)
    for c in candidate_counts(group,vendors).values():assert not c['串码'] and not c['换货']
    meta={'both_files':True,'coverage':True,'evidence':True,'sales_cost':10000}
    assert not any('待核实' in x for x in export_blockers(group,meta,vendors))

def test_supplier_append_and_conflict(items):
    bc=items[0]['商品条码']
    v=[{'商品条码':bc,'供货商名称':'新增供应商'}]
    assert supplier_map(v)[bc]=='新增供应商'
    v.append({'商品条码':bc,'供货商名称':'第二供应商'})
    assert supplier_map(v)[bc]=='多供货商·待核实'
    assert default_vendors()==[]

@pytest.mark.parametrize('value',['NaN','Infinity','-Infinity',True,'错误',-1])
def test_invalid_counts_cannot_be_saved(items,value):
    with pytest.raises(ValueError):patch_items(items,{items[0]['记录ID']:{'复盘实数':value}})

def ready(items):
    for x in items:
        x['复盘实数']=x['初盘数'];x['备注']='现场已核实';x['核实结论']='已排除'
    net=summary(items,None)['净差异']
    return {'both_files':True,'coverage':True,'evidence':True,'sales_cost':-net*500}

def test_dynamic_signed_rate_boundary_and_completeness(items):
    meta=ready(items)
    assert summary(items,meta['sales_cost'])['报损率']==-2
    assert not export_blockers(items,meta,default_vendors())
    meta['sales_cost']*=D('.9')
    assert any('超过2' in x for x in export_blockers(items,meta,default_vendors()))
    meta['sales_cost']*=2;meta['coverage']=False
    assert any('完整' in x for x in export_blockers(items,meta,default_vendors()))
    meta['coverage']=True;items[0]['备注']='123'
    assert any('原因' in x for x in export_blockers(items,meta,default_vendors()))
    items[0]['备注']='已核实';items[0]['原始调整']=D('.4')
    assert any('特殊单位' in x for x in export_blockers(items,meta,default_vendors()))

def test_company_file_has_exactly_three_typed_columns(items):
    meta=ready(items)
    blob=create_company_xls(items,meta,default_vendors())
    book=xlrd.open_workbook(file_contents=blob);s=book.sheet_by_index(0)
    assert s.ncols==3
    assert tuple(s.row_values(0))==('商品编码','商品条码','数量')
    assert s.nrows==len(items)+1
    for n,x in enumerate(items,1):
        assert s.cell_type(n,0)==s.cell_type(n,1)==xlrd.XL_CELL_TEXT
        assert s.cell_type(n,2)==xlrd.XL_CELL_NUMBER
        assert s.row_values(n)==[x['商品编码'],x['商品条码'],float(current_total(x))]
    items[0]['复盘实数']=None
    with pytest.raises(ValueError):create_company_xls(items,meta,default_vendors())

def test_duplicates_never_overwrite_or_sum(items):
    items[1]['商品条码']=items[0]['商品条码']
    assert any('重复' in x for x in validate_all(items))
    items[1]['商品条码']='unique';items[1]['商品编码']=items[0]['商品编码']
    assert any('重复' in x for x in validate_all(items))

def test_negative_system_valid_but_missing_cost_not_zero(items):
    items[0]['系统数']=D(-1)
    assert not validate_all(items)
    items[0]['成本价']=None
    assert validate_all(items)
    assert summary(items,100)['净差异'] is None
    assert summary(items,100)['报损率'] is None

def test_v13_sync_star_headers_zero_and_duplicate_atomic(items):
    wb=openpyxl.Workbook();s=wb.active;s.title='01现场复盘'
    s.append(['商品条码','商品编码','复盘实数★','备注／原因★'])
    x=items[0];s.append([x['商品条码'],x['商品编码'],0,'现场为零'])
    out=io.BytesIO();wb.save(out)
    updated,count=import_v13(out.getvalue(),'复盘.xlsx',items)
    assert count==1 and updated[0]['复盘实数']==0 and items[0]['复盘实数'] is None
    s.append([x['商品条码'],x['商品编码'],1,'重复'])
    out=io.BytesIO();wb.save(out)
    with pytest.raises(ValueError):import_v13(out.getvalue(),'复盘.xlsx',items)
    assert items[0]['复盘实数'] is None

def test_all_sheets_and_barcode_precision(sample):
    name,blob=sample[3]['盘盈'];source=xlrd.open_workbook(file_contents=blob).sheet_by_index(0)
    wb=xlwt.Workbook()
    for title in ('第一表','第二表'):
        s=wb.add_sheet(title)
        for r in range(2):
            for c,v in enumerate(source.row_values(r)):s.write(r,c,v)
    out=io.BytesIO();wb.save(out)
    rows,_=read_source(out.getvalue(),'多表.xls','盘盈')
    assert len(rows)==2
    assert rows[0]['工作表']!=rows[1]['工作表']
    # 格式显示前导零不可猜测；禁止隐式截断长条码。
    wb=xlwt.Workbook();s=wb.add_sheet('源表')
    for c,v in enumerate(source.row_values(0)):s.write(0,c,v)
    for c,v in enumerate(source.row_values(1)):
        if c==2:s.write(1,c,123,xlwt.easyxf(num_format_str='000000'))
        else:s.write(1,c,v)
    out=io.BytesIO();wb.save(out)
    rows,_=read_source(out.getvalue(),'精度.xls','盘盈')
    assert any(p['等级']=='阻断' and p['规则']=='ID_LEADING_ZERO' for p in rows[0]['导入问题'])
