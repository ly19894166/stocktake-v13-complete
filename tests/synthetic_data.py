"""从零构造虚构样本；不读取正式目录、公司文件或真实供应商目录。"""
import io
import sys
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import xlwt
from source_core import INPUT_HEADERS

def source_rows(count=12):
    rows=[]
    for i in range(count):
        system=20
        difference=(i%5+2) if i%2==0 else -(i%5+7)
        if i==count-1:difference=0
        cost=Decimal('3.50')+Decimal(i%3)
        rows.append(dict(zip(INPUT_HEADERS,[f'SIM{i:05d}',f'SHOP{i:05d}',f'099{i:010d}',
            f'模拟商品{i:03d}','件','模拟规格','模拟食品部','模拟食品类','模拟子类','模拟细类',
            system,system+difference,0,None,difference,0,float(cost),
            float((Decimal(difference)*cost).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)),''])))
    # 故意生成一个旧报差异，与初盘/系统数不符；验证不会偷用旧值修正库存。
    if count>2:rows[2]['差异数量']+=3
    return rows

def workbook(rows):
    book=xlwt.Workbook();sheet=book.add_sheet('模拟初盘')
    for c,h in enumerate(INPUT_HEADERS):sheet.write(0,c,h)
    for r,row in enumerate(rows,1):
        for c,h in enumerate(INPUT_HEADERS):
            v=row[h]
            if v is not None:sheet.write(r,c,v)
    out=io.BytesIO();book.save(out);return out.getvalue()

def simulated_pair(count=12):
    rows=source_rows(count)
    gain=[r for r in rows if r['初盘数量']>=r['库存数量']]
    loss=[r for r in rows if r['初盘数量']<r['库存数量']]
    return ('模拟盘盈.xls',workbook(gain)),('模拟盘亏.xls',workbook(loss))

def simulated_master(count=37):
    import openpyxl
    book=openpyxl.Workbook();sheet=book.active;sheet.title='01现场复盘'
    archive=book.create_sheet('04商品档案后台');vendor=book.create_sheet('05供货商目录')
    for i,row in enumerate(source_rows(count),6):
        for c,value in {1:i-5,2:row['大类名称'],3:row['商品名称'],4:row['商品条码'],
                        5:row['商品'],6:row['库存数量'],7:row['初盘数量']}.items():sheet.cell(i,c,value)
        for c,value in {2:row['商品'],3:row['商品商家编码'],7:row['部类名称'],8:row['大类名称'],
                        9:row['中类名称'],10:row['小类名称'],11:row['成本价'],12:8.5,
                        13:0.4 if i==6 else 0}.items():archive.cell(i,c,value)
    vendor.cell(6,1,'0990000000000');vendor.cell(6,2,'SIM-VENDOR');vendor.cell(6,3,'模拟供货商甲')
    out=io.BytesIO();book.save(out);return out.getvalue()

def write_examples(directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    for name,blob in simulated_pair():(directory/name).write_bytes(blob)
    (directory/'模拟V13母版.xlsx').write_bytes(simulated_master())
    return directory

if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser(description='仅生成虚构商品数据，输出文件不提交Git')
    parser.add_argument('--output',default='tests/runtime/simulated')
    print(write_examples(parser.parse_args().output))
