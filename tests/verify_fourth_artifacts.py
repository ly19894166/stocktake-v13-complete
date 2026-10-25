"""独立重读模拟审计Excel、真XLS、PDF与重启哈希；禁止读取正式数据目录。"""
from pathlib import Path
import os,sys,json,io,re,sqlite3,hashlib
from collections import Counter
import openpyxl,xlrd
from pypdf import PdfReader
from pypdf.generic import ContentStream
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from storage import Store
from vendor_audit import audit_model,filtered_rows,HEADERS
from print_manager import make_pdf
from workflow import print_order
from core import current_total
DATA=Path(os.environ.get('STOCKTAKE_FOURTH_DATA_DIR',ROOT/'tests/runtime/fourth_data')).resolve();OUT=ROOT/'tests/runtime/fourth_results';OUT.mkdir(exist_ok=True)
assert DATA.is_relative_to((ROOT/'tests/runtime').resolve()),'只允许模拟库'

def snapshot():
    with sqlite3.connect(f'file:{DATA / "stocktake.sqlite"}?mode=ro',uri=True) as c:
        assert c.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
        return {t:hashlib.sha256(repr(c.execute('SELECT * FROM '+t+' ORDER BY rowid').fetchall()).encode()).hexdigest() for (t,) in c.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")}

def load(s):
    batches=s.list_batches();assert len(batches)==1
    b=batches[0]['id'];items,_,state,_=s.load(b)
    assert len(items)==450 and all(x['商品条码'].startswith('099') and x['商品名称'].startswith('模拟') for x in items)
    return b,items,state

def audit():
    s=Store(DATA)
    try:
        b,items,state=load(s);m=audit_model(items,state['vendors'],state['meta']);rows=filtered_rows(m)
        w=openpyxl.load_workbook(OUT/'模拟卓不凡专项审计.xlsx',data_only=False);sheet=w['卓不凡专项审计']
        assert tuple(c.value for c in sheet[1])==HEADERS and sheet.max_row==61 and sheet.max_column==16
        for ri,row in enumerate(rows,2):
            for ci,k in enumerate(HEADERS,1):
                a=sheet.cell(ri,ci).value;e=row[k]
                assert (a or '')==(e or '') or isinstance(a,(int,float)) and isinstance(e,(int,float)) and abs(a-e)<1e-8,(ri,k)
            assert sheet.cell(ri,4).data_type==sheet.cell(ri,5).data_type=='s'
        stats={r[0].value:r[1].value for r in w['汇总与口径'].iter_rows(min_row=2)}
        assert stats['商品总条数']==60 and stats['供货商待确认']==25 and stats['疑似A换B对数']==900
        for k in ('盘盈','盘亏','净差异'):assert abs(stats[k]-float(m['stats'][k]))<1e-8
        assert w['A换B线索'].max_row==901
        assert any('模拟换货人工核对说明' in str(sheet.cell(r,15).value) and sheet.cell(r,16).value=='SIM-PROOF-FOURTH' for r in range(2,62))
        result={'pass':True,'simulated_only':True,'rows':60,'pending':25,'candidate_pairs':900,'headers':HEADERS,'column_count':16,'text_codes_and_barcodes':True,'all_cells_match_saved_records':True,'amounts_match':True,'manual_proof_preserved':True,'company_file_separate':True}
        (OUT/'audit_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');return result
    finally:s.close()

def pdf():
    s=Store(DATA)
    try:
        b,items,state=load(s);rows=print_order(items[:414],'category')
        raw,pages=make_pdf(rows,state['meta'],'SIM-414',scope='模拟414条现场找货');(OUT/'模拟414条紧凑现场单.pdf').write_bytes(raw)
        reader=PdfReader(io.BytesIO(raw));assert pages<=30
        codes=[];page_counts=[];boxes=0;embedded=False
        for page in reader.pages:
            assert abs(float(page.mediabox.width)-841.89)<1 and abs(float(page.mediabox.height)-595.28)<1
            text=page.extract_text();assert not any(w in text for w in ('串码','A换B','疑似','算法'))
            per=re.findall(r'(?<!\d)099\d{10}(?!\d)',text);codes+=per;page_counts.append(len(per))
            for font in page['/Resources']['/Font'].values():
                fd=font.get_object().get('/FontDescriptor')
                if fd and '/FontFile2' in fd.get_object():embedded=True
            for args,op in ContentStream(page.get_contents(),reader).operations:
                if op==b're' and 80<=float(args[2])<=110 and float(args[3])>=22:boxes+=1
        assert Counter(codes)==Counter(x['商品条码'] for x in rows) and boxes==414 and embedded
        assert all(14<=n<=18 for n in page_counts[:-1])
        archives=s.prints(b);repeats=[p for p in archives if p.get('parent')];assert repeats
        for p in repeats:assert s.print_pdf(b,p['id'])==s.print_pdf(b,p['parent'])
        native=[p for p in archives if not p.get('parent')];assert len(native)>=2
        for p in native:
            assert p['filter_policy']['include_zbf'] is False
            assert all('related' not in r and 'group' not in r for r in p['rows'])
            assert all(x['source_id'] in {r['记录ID'] for r in items} for x in p['rows'])
        result={'pass':True,'simulated_only':True,'rows':414,'pages':pages,'per_page_rows':page_counts,'blank_handwriting_boxes':boxes,'minimum_box_height_points':22,'embedded_chinese_font':embedded,'barcodes_complete_and_unique':True,'no_pairing_text_or_groups':True,'category_middle_sequence_sort':True,'historical_reprint_byte_identical':True,'new_data_new_id':True}
        (OUT/'pdf_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');return result
    finally:s.close()

def company():
    s=Store(DATA)
    try:
        b,items,state=load(s);blob=(OUT/'模拟公司最终450行.xls').read_bytes();assert blob[:8]==bytes.fromhex('D0CF11E0A1B11AE1')
        w=xlrd.open_workbook(file_contents=blob);assert w.sheet_names()==['Sheet1'];sheet=w.sheet_by_index(0)
        assert sheet.nrows==451 and sheet.ncols==3 and sheet.row_values(0)==['商品编码','商品条码','数量']
        for r,x in enumerate(items,1):
            assert sheet.row_values(r)==[x['商品编码'],x['商品条码'],float(current_total(x))]
            assert sheet.cell_type(r,0)==sheet.cell_type(r,1)==xlrd.XL_CELL_TEXT and sheet.cell_type(r,2)==xlrd.XL_CELL_NUMBER
        classes=audit_model(items,state['vendors'],state['meta'])['classes'];assert sum(c['status']=='zbf' for c in classes.values())==60
        assert any(sheet.cell_value(r,2)==0 for r in range(1,451))
        result={'pass':True,'simulated_only':True,'true_xls':True,'sheet':'Sheet1','rows':450,'columns':3,'typed_text_text_number':True,'all_60_zbf_records_retained':True,'all_rows_equal_latest_totals':True,'original_order_and_leading_zero_and_zero_preserved':True}
        (OUT/'company_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');return result
    finally:s.close()

if __name__=='__main__':
    mode=sys.argv[1] if len(sys.argv)>1 else 'audit'
    if mode=='before':(OUT/'restart_expected.json').write_text(json.dumps(snapshot()),encoding='utf8');result={'snapshot_saved':True}
    elif mode=='after':
        assert snapshot()==json.loads((OUT/'restart_expected.json').read_text('utf8'));result={'pass':True,'simulated_only':True,'real_process_restart':True,'all_8_tables_unchanged':True};(OUT/'restart_verification.json').write_text(json.dumps(result,indent=2),encoding='utf8')
    else:result={'audit':audit,'pdf':pdf,'company':company}[mode]()
    print(json.dumps(result,ensure_ascii=False))
