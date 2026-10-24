"""可选证据、供应商通用表头、V13完整工作簿与明确标识草稿。"""
import copy
import hashlib
import io
import json
import re
from collections import defaultdict
from decimal import Decimal as D
from pathlib import Path
from core import (load_excel,as_code,decimal_value,hydrate,rounded,VENDOR_HEADERS,
    current_total,newest_diff,current_amount,changes,summary,change_rows,validate_all,VERSION)

def header(v):return str(v or '').strip().replace('★','').replace('：',':')

def identify_batch(sources):
    meta={}
    for name,blob in sources.values():
        for _,rows in load_excel(blob,name):
            for row in rows[:20]:
                for i,v in enumerate(row):
                    text=header(v)
                    for key,labels in [('store',('门店名称','门店','店仓名称')),('period',('盘点期间','盘点日期'))]:
                        known_headers={v for aliases in ALIASES.values() for v in aliases}
                        if text in labels and i+1<len(row) and row[i+1] not in (None,'') and sum(v not in (None,'') for v in row)<=4 and header(row[i+1]) not in known_headers:
                            meta.setdefault(key,str(row[i+1]))
                        for label in labels:
                            if text.startswith(label+':') and text.split(':',1)[1].strip():
                                meta.setdefault(key,text.split(':',1)[1].strip())
            # 明确的门店/期间数据列：只有列内标识唯一时才自动识别。
            for i,row in enumerate(rows[:20]):
                heads=[header(v) for v in row]
                if '商品条码' not in heads:continue
                for key,labels in [('store',('门店名称','门店','店仓名称')),('period',('盘点期间','盘点日期'))]:
                    for label in labels:
                        if label in heads:
                            c=heads.index(label);values={str(r[c]).strip() for r in rows[i+1:] if c<len(r) and r[c] not in (None,'')}
                            if len(values)==1:meta.setdefault(key,next(iter(values)))
                break
    return meta

ALIASES={'商品条码':('商品条码','条码','条形码','主条码','国际条码','商品主条码'),
    '供货商名称':('供货商名称','供应商名称','供货商','供应商','厂商'),
    '供货商编码':('供货商编码','供应商编码'), '商品编码':('商品编码','商品','商品代码'),
    '商品名称':('商品名称','品名','商品名'),'数量':('数量','销售数量','销售数','出入库数量','实发数量','收货数量'),
    '差异数量':('最新差异','复核差异','差异数量','差异数','盘盈盘亏数量'),
    '系统数':('库存数量','系统数','系统库存','库存'),
    '时间':('业务日期','日期','单据日期','销售日期','时间','发生时间'),
    '单号':('单据号','单号','单据编号','订单号'),'包装系数':('包装系数','箱规','件装数')}

def generic_rows(blob,name,required=('商品条码',),mapping=None):
    output=[];matched=False
    from source_core import read_workbook
    books=read_workbook(blob,name)
    metadata={s['name']:s['cells'] for s in books}
    sheets=[(s['name'],s['rows']) for s in books]
    if name.lower().endswith('.xls'):
        import xlrd
        book=xlrd.open_workbook(file_contents=blob)
        sheets=[(s.name,[[xlrd.xldate.xldate_as_datetime(c.value,book.datemode).isoformat(sep=' ') if c.ctype==xlrd.XL_CELL_DATE else c.value for c in s.row(r)] for r in range(s.nrows)]) for s in book.sheets()]
        book.release_resources()
    for sheet,rows in sheets:
        found=None
        for i,row in enumerate(rows[:30]):
            heads=[header(v) for v in row];indexes={}
            for key,aliases in ALIASES.items():
                chosen=(mapping or {}).get(key)
                opts=(chosen,) if chosen else aliases
                for alias in opts:
                    if alias in heads:indexes[key]=heads.index(alias);break
            if all(k in indexes for k in required):found=(i,heads,indexes);break
        if found is None:continue
        matched=True;start,heads,indexes=found
        for rowno,row in enumerate(rows[start+1:],start+2):
            def val(c):return row[c] if c<len(row) else None
            if not any(v not in (None,'') for v in row):continue
            d={k:val(c) for k,c in indexes.items()}
            if any(as_code(d.get(k)) in ('合计','总计','小计') for k in ('商品条码','商品名称','商品编码')):continue
            if not as_code(d.get('商品条码')):raise ValueError(f'{sheet} 第{rowno}行缺条码，未导入')
            raw=d['商品条码']
            if isinstance(raw,(int,float)) and len(as_code(raw))>15:raise ValueError('数值条码超过15位，须使用完整文本原值')
            if isinstance(raw,(int,float)):
                fmt=metadata[sheet][rowno-1][indexes['商品条码']].get('number_format','')
                if re.fullmatch(r'0+',fmt) and len(fmt)>len(as_code(raw)):
                    raise ValueError('条码依赖数字格式显示前导零，须核对并以完整文本条码导入')
            if isinstance(raw,str) and re.fullmatch(r'[+-]?\d+(\.\d+)?[eE][+-]?\d+',raw.strip()):
                raise ValueError('条码为科学计数法，不能猜测完整条码')
            d['商品条码']=as_code(raw)
            d.update({'工作表':sheet,'源行号':rowno,'原始资料':{h:val(c) for c,h in enumerate(heads) if h}})
            for key in ('数量','差异数量','系统数','包装系数'):
                if d.get(key) not in (None,''):
                    value=decimal_value(d[key])
                    if value is None:raise ValueError(f'{sheet} 第{rowno}行{key}无效')
                    d[key]=str(value)
            output.append(d)
    if not matched:raise ValueError('未识别表头，请检查或选择列映射：'+'、'.join(required))
    if not output:raise ValueError('没有可导入明细')
    return output

def supplier_import(blob,name,mapping=None):
    rows=generic_rows(blob,name,('商品条码','供货商名称'),mapping)
    result=[]
    for row in rows:
        raw=row['原始资料'];v={k:raw.get(k) for k in VENDOR_HEADERS}
        v.update({k:as_code(row.get(k)) for k in ('商品条码','供货商编码','供货商名称')})
        if not v['供货商名称']:raise ValueError('供应商名称不能为空')
        v['供应商商品名称']=raw.get('供应商商品名称') or row.get('商品名称') or ''
        for key in ('包装系数',):
            if row.get(key):v[key]=row[key]
        codes=[v['商品条码']]
        for key,value in raw.items():
            if value not in ('',None) and re.search(r'副条码|条码\d|多条码|辅助条码|主条码',key):
                codes.extend(c for c in re.split(r'[;；,，、\s]+',as_code(value)) if c)
        for code in dict.fromkeys(codes):result.append({**v,'商品条码':code})
    return result

def cross_candidates(items,evidence):
    bybarcode=defaultdict(list)
    for file in evidence:
        if file['类型']!='跨店差异':continue
        for r in file['rows']:
            d=decimal_value(r.get('差异数量'))
            if d not in (None,0):bybarcode[r['商品条码']].append((file,r,d))
    result=defaultdict(list)
    for x in items:
        diff=newest_diff(x)
        if diff in (None,0):continue
        for file,row,other in bybarcode[x['商品条码']]:
            if diff*other<0:
                result[x['记录ID']].append({'门店':file['门店'],'条码':x['商品条码'],
                    '本店差异':str(diff),'对方原报差异':str(other),'单号':str(row.get('单号') or ''),
                    '日期':str(row.get('时间') or '未提供时点·需核同期'),'文件':file['文件'],'源行号':row['源行号'],
                    '提示':'疑似跨店调拨·仅作证据，不合并实数、不自动平账'})
    return dict(result)

def transactions_for(barcode,evidence,start='',end=''):
    result=[]
    for file in evidence:
        if file['类型'] not in ('销售明细204','调拨/入库/报损单','最新库存','跨店库存'):continue
        for r in file['rows']:
            if r['商品条码']!=barcode:continue
            at=str(r.get('时间') or '')
            if at and ((start and at<start) or (end and at[:10]>end)):continue
            result.append({'证据类型':file['类型'],'商品条码':barcode,'数量':r.get('数量'),
                '原报库存':r.get('系统数'),'日期':at or '未提供时点·不可认定为盘点间交易',
                '单号':str(r.get('单号') or ''),'文件':file['文件'],'源行号':r['源行号']})
    return result

def import_master(blob,name):
    from openpyxl import load_workbook
    from source_core import INPUT_HEADERS
    w=load_workbook(io.BytesIO(blob),data_only=False)
    if not all(s in w.sheetnames for s in ('01现场复盘','04商品档案后台')):raise ValueError('不是完整V13母版')
    s=w['01现场复盘'];a=w['04商品档案后台'];sha=hashlib.sha256(blob).hexdigest();items=[]
    for r in range(6,s.max_row+1):
        if not s.cell(r,4).value or isinstance(s.cell(r,1).value,str) and s.cell(r,1).value.startswith('='):continue
        v=lambda col:s.cell(r,col).value
        av=lambda col:a.cell(r,col).value
        system=decimal_value(v(6));initial=decimal_value(v(7));adj=decimal_value(av(13),D(0))
        if system is None or initial is None:raise ValueError(f'V13第{r}行系统数/初盘无效')
        cost=decimal_value(av(11));diff=rounded(initial-system+adj)
        raw=[as_code(av(2)),as_code(av(3)),as_code(v(4)),str(v(3) or ''),'','',str(av(7) or ''),str(av(8) or ''),str(av(9) or ''),str(av(10) or ''),
             float(system),float(initial),float(adj),None,float(diff),None,None if cost is None else float(cost),None if cost is None else float(rounded(diff*cost,2)),'']
        item={'记录ID':sha[:12]+':01现场复盘:'+str(r),'文件SHA256':sha,'来源':'V13母版','原文件':name,
            '工作表':'01现场复盘','原始行号':r,'序号':len(items)+1,'商品编码':as_code(av(2)),
            '商家编码':as_code(av(3)),'商品条码':as_code(v(4)),'商品名称':str(v(3) or ''),
            '部类':str(av(7) or ''),'大类':str(av(8) or v(2) or ''),'中类':str(av(9) or ''),'小类':str(av(10) or ''),
            '系统数':system,'初盘数':initial,'原始调整':adj,'子转母':adj,'成本价':cost,
            '源报差异数':diff,'原始差异成本':None if cost is None else rounded(diff*cost,2),
            '复盘实数':decimal_value(v(8)),'备注':str(v(11) or ''),'单位说明':'','零售价':decimal_value(av(12)),
            '核实结论':'','相关单号':'','导入问题':[],'原始字段':dict(zip(INPUT_HEADERS,raw)),
            '原始整行':raw,'件装单位':'','销售规格':'','品牌':'','包装系数':None,'源表备注':''}
        items.append(item)
    if not items:raise ValueError('V13没有商品明细')
    vendors=[]
    if '05供货商目录' in w:
        for row in w['05供货商目录'].iter_rows(min_row=6,values_only=True):
            if row[0]:vendors.append(dict(zip(VENDOR_HEADERS,row[:len(VENDOR_HEADERS)])))
    w.close()
    return items,[],sha[:20],vendors

def export_workbook(items,state,engine,draft=True,candidate_scope='all'):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill,Font,Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook();wb.remove(wb.active)
    pair_count=sum(engine.count(x['记录ID'],'all') for x in items)//2
    if pair_count>1048573 and candidate_scope=='all':
        raise ValueError(f'成对候选约{pair_count}对，超过单张Excel容量；请选择候选摘要。软件内仍可完整分页查看，不改变任何盘点数量。')
    banner='草稿·未批准·禁止作为公司库存调整提交' if draft else '审批准备资料·仍须人工按公司流程审核'
    def table(title,columns,rows):
        s=wb.create_sheet(title);s.append([banner]);s.append(columns)
        for row in rows:
            s.append([float(v) if isinstance(v,D) else v for v in row])
        s.freeze_panes='F3';s.auto_filter.ref=f'A2:{get_column_letter(len(columns))}{max(2,s.max_row)}'
        s.sheet_properties.pageSetUpPr.fitToPage=True
        s.print_title_rows='1:2';s.oddHeader.center.text=banner;s.sheet_properties.tabColor='087F8C'
        for ce in s[1]:ce.font=Font(color='B42318',bold=True,size=13)
        for ce in s[2]:ce.fill=PatternFill('solid',fgColor='19364C');ce.font=Font(color='FFFFFF',bold=True)
        for col in range(1,len(columns)+1):s.column_dimensions[get_column_letter(col)].width=22
        for row in s.iter_rows(min_row=3):
            for ce,col in zip(row,columns):
                if any(k in col for k in ('条码','编码')):ce.number_format='@';ce.value='' if ce.value is None else str(ce.value)
                if isinstance(ce.value,str) and ce.value.startswith(('=','+','-','@')):ce.data_type='s'
        return s
    def tag(x):
        peers=engine.peers(x['记录ID'])
        return '\n'.join(f"{p['kind']} ↔ {p['name']} / {p['barcode']} / 差异{p['diff']:+g} / {p['confidence']}" for p in peers)
    table('01现场复盘',['序号','品类','商品名称','商品条码','商品编码','系统数','初盘数','复盘实数','当前实物总数','最新差异','备注／原因','供货商','核查标签'],
       [[x['序号'],x['大类'],x['商品名称'],x['商品条码'],x['商品编码'],x['系统数'],x['初盘数'],x.get('复盘实数'),current_total(x),newest_diff(x),x.get('备注',''),engine.suppliers.get(x['商品条码'],'未在目录'),tag(x)] for x in items])
    rows=change_rows(items,state['vendors'])
    for row,x in zip(rows,items):
        before=x['初盘数']-x['系统数']+x['原始调整'] if x['初盘数'] is not None and x['系统数'] is not None else None
        after=newest_diff(x)
        row.update({'盘亏扩大量':None if before is None or after is None else float(max(D(0),max(D(0),-after)-max(D(0),-before))),
            '跨正负变化':'是' if before is not None and after is not None and before*after<0 else '否',
            '核查标签':tag(x),'原始文件':x['原文件'],'原始行号':x['原始行号']})
    cols=list(rows[0]) if rows else []
    table('02自动清单与系统复制',cols,[[r.get(c) for c in cols] for r in rows])
    stats=summary(items,state['meta'].get('sales_cost'))
    table('03汇总与报损审核',['项目','当前值'],list(stats.items())+[['2‰净损耗标准',2],['完整范围确认',state['meta'].get('coverage',False)],['凭证确认',state['meta'].get('evidence',False)]])
    cols=['序号','商品编码','商家编码','商品条码','商品名称','部类','大类','中类','小类','成本价','零售价','原始调整','单位说明','包装系数','原文件','原始行号']
    table('04商品档案后台',cols+['初盘差异','最新差异','差异成本','串码候选数量','A换B候选数量','输入错误'],
        [[x.get(k) for k in cols]+[x['初盘数']-x['系统数']+x['原始调整'] if x['初盘数'] is not None and x['系统数'] is not None else None,
        newest_diff(x),current_amount(x),len(engine.neighbors(x['记录ID'])),len(engine.neighbors(x['记录ID'],'swap')),
        '；'.join(p['说明'] for p in x.get('导入问题',[]) if p['等级']=='阻断')] for x in items])
    table('05供货商目录',list(VENDOR_HEADERS),[[v.get(k) for k in VENDOR_HEADERS] for v in state['vendors']])
    pairs=[];seen=set()
    for x in items:
        for p in engine.peers(x['记录ID'],limit=3 if candidate_scope=='summary' else engine.count(x['记录ID'],'all')):
            key=tuple(sorted((p['a'],p['b'])))
            if key in seen:continue
            seen.add(key);a,b=engine.items[p['a']],engine.items[p['b']]
            pairs.append([a['序号'],a['商品名称'],a['商品条码'],str(engine.diff[p['a']]),b['序号'],b['商品名称'],b['商品条码'],str(engine.diff[p['b']]),p['kind'],p['confidence'],p['evidence']])
    table('06异常候选后台',['A序号','A商品','A条码','A最新差异','B序号','B商品','B条码','B最新差异','候选类型','线索强度','证据'],pairs)
    table('07参数与说明',['参数','值'],[['版本',VERSION],['候选规则','成本按分相同，任一分类层级相同，反向差异；零差异撤销'],
       ['异常候选导出范围','每商品前3条摘要；完整候选在软件分页可查' if candidate_scope=='summary' else '完整成对候选'],
       ['原始来源','；'.join(state['meta'].get('source_files',[]))],['首次导入北京时间',state['meta'].get('first_import_at','')],
       ['门店',state['meta']['store']],['期间',state['meta']['period']]]+[[k,str(v)] for k,v in engine.parameters.items()])
    out=io.BytesIO();wb.save(out);return out.getvalue()
