"""V1.1 import kernel: immutable source evidence, explicit exclusions and recounts."""
from __future__ import annotations
import copy
import csv
import hashlib
import io
import json
import math
import os
import re
import zipfile
from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

VERSION='1.1.0'
INPUT_HEADERS=('商品','商品商家编码','商品条码','商品名称','件装单位','销售规格',
 '部类名称','大类名称','中类名称','小类名称','库存数量','初盘数量',
 '子转母商品数量','复盘数量','差异数量','销售数量','成本价','成本金额','初复盘差异原因解释')
SOURCE_REQUIRED=('商品','商品条码','商品名称','库存数量','初盘数量','差异数量','成本价','成本金额')
NUMERIC=('库存数量','初盘数量','子转母商品数量','复盘数量','差异数量','销售数量','成本价','成本金额')
OUTPUT_HEADERS=('商品编码','商品条码','数量')
ALIASES={'商品编码':'商品','系统库存':'库存数量','初盘数':'初盘数量','差异数':'差异数量'}
EDIT_FIELDS=('复盘实数','备注','单位修正量','单位说明','处理','处理说明')

def decimal_value(value,default=None):
    if value is None or isinstance(value,bool):return default
    text=str(value).strip().replace('−','-').replace('－','-')
    if not text:return default
    if text.startswith('(') and text.endswith(')'):text='-'+text[1:-1]
    if ',' in text:
        if not re.fullmatch(r'[+-]?\d{1,3}(,\d{3})+(\.\d+)?',text):return default
        text=text.replace(',','')
    try:
        result=Decimal(text)
        return result if result.is_finite() else default
    except (InvalidOperation,ValueError):return default

def num(value):
    d=decimal_value(value)
    return None if d is None else float(d)

def as_code(value):
    if value is None or isinstance(value,bool):return ''
    if isinstance(value,(int,float)):
        return str(int(value)) if math.isfinite(value) and value==int(value) else ''
    return str(value).strip()

def header_name(value):
    text=str(value or '').replace('\ufeff','').replace('\u3000',' ').strip()
    return ALIASES.get(text,text)

def is_red(rgb):return bool(rgb and len(rgb)>=3 and rgb[0]>=150 and rgb[1]<100 and rgb[2]<100)

def colname(c):
    out=''
    while c>=0:out=chr(c%26+65)+out;c=c//26-1
    return out

def read_workbook(blob,filename):
    sheets=[];suffix=os.path.splitext(filename.lower())[1]
    if suffix=='.xls':
        import xlrd
        book=xlrd.open_workbook(file_contents=blob,formatting_info=True)
        try:
            for sh in book.sheets():
                rows=[];metadata=[]
                for r in range(sh.nrows):
                    values=[];cells=[]
                    for c in range(sh.ncols):
                        ce=sh.cell(r,c);xf=book.xf_list[ce.xf_index];font=book.font_list[xf.font_index]
                        rgb=book.colour_map.get(font.colour_index);fill=book.colour_map.get(xf.background.pattern_colour_index)
                        values.append(ce.value)
                        cells.append({'cell':f'{colname(c)}{r+1}','type':ce.ctype,'xf':ce.xf_index,
                            'font_rgb':rgb,'fill_rgb':fill,'number_format':book.format_map[xf.format_key].format_str,
                            'red_font':is_red(rgb),'red_fill':is_red(fill) and xf.background.fill_pattern!=0,
                            'error':ce.ctype==xlrd.XL_CELL_ERROR})
                    rows.append(values);metadata.append(cells)
                sheets.append({'name':sh.name,'rows':rows,'cells':metadata,'visibility':sh.visibility})
        finally:book.release_resources()
    elif suffix=='.xlsx':
        from openpyxl import load_workbook
        from openpyxl.styles.colors import COLOR_INDEX as COLOR_INDEXED
        book=load_workbook(io.BytesIO(blob),data_only=True)
        formulas=load_workbook(io.BytesIO(blob),data_only=False)
        def rgb(color):
            if color is None:return None
            text=color.rgb if color.type=='rgb' else COLOR_INDEXED[color.indexed] if color.type=='indexed' and color.indexed<len(COLOR_INDEXED) else None
            if not isinstance(text,str):return None
            try:return tuple(int(text[-6:][i:i+2],16) for i in (0,2,4))
            except ValueError:return None
        try:
            for sh in book:
                rows=[];metadata=[]
                for row in sh:
                    values=[];cells=[]
                    for ce in row:
                        fo=formulas[sh.title][ce.coordinate];frgb=rgb(ce.font.color);fill=rgb(ce.fill.fgColor)
                        value=ce.value
                        if value is not None and not isinstance(value,(str,int,float,bool)):value=str(value)
                        values.append(value)
                        cells.append({'cell':ce.coordinate,'type':ce.data_type,'font_rgb':frgb,'fill_rgb':fill,
                            'number_format':ce.number_format,'red_font':is_red(frgb),
                            'red_fill':is_red(fill) and ce.fill.patternType=='solid','formula':fo.value if fo.data_type=='f' else None,
                            'error':ce.data_type=='e'})
                    rows.append(values);metadata.append(cells)
                sheets.append({'name':sh.title,'rows':rows,'cells':metadata,'visibility':sh.sheet_state})
        finally:book.close();formulas.close()
    else:raise ValueError('仅支持真正的 .xls / .xlsx；不能只更改扩展名。')
    return sheets

def load_excel(blob,filename):return [(s['name'],s['rows']) for s in read_workbook(blob,filename)]

def locate_source_sheet(sheets):
    for name,rows in sheets:
        for ix,row in enumerate(rows[:20]):
            head=[header_name(v) for v in row]
            if all(h in head for h in SOURCE_REQUIRED):return name,ix,head,rows
    raise ValueError('未找到必需字段：'+'、'.join(SOURCE_REQUIRED))

def location(item):return f"{item['原文件']} / {item['工作表']} / 第{item['原始行号']}行"

def issue(code,level,message,item=None):
    return {'规则':code,'等级':level,'位置':location(item) if item else '',
            '记录ID':item['记录ID'] if item else '', '说明':message}

def read_source(blob,filename,expected,exclude_red=False):
    if expected not in ('盘盈','盘亏'):raise ValueError('来源槽位须为盘盈或盘亏')
    digest=hashlib.sha256(blob).hexdigest();items=[];notes=[];matched=0
    for sh in read_workbook(blob,filename):
        found=None
        for ix,row in enumerate(sh['rows'][:20]):
            head=[header_name(v) for v in row]
            if all(h in head for h in SOURCE_REQUIRED):found=(ix,head);break
        if found is None:
            notes.append(issue('SHEET_SKIPPED','注意',f"{filename} / {sh['name']}未找到标准表头，请核对是否有未识别明细"));continue
        matched+=1;idx,headers=found
        dup=[h for h in set(headers) if h and headers.count(h)>1]
        if dup:raise ValueError(f"{sh['name']}存在重名表头：{dup}")
        lookup={h:c for c,h in enumerate(headers) if h}
        for r in range(idx+1,len(sh['rows'])):
            vals=sh['rows'][r];cells=sh['cells'][r]
            if not any(v is not None and str(v).strip() for v in vals):continue
            if all(header_name(vals[lookup[h]])==h for h in SOURCE_REQUIRED):continue
            def v(h):
                c=lookup.get(h);return vals[c] if c is not None and c<len(vals) else None
            if any(as_code(v(h)) in ('合计','小计','总计') for h in ('商品','商品条码','商品名称')):
                notes.append(issue('TOTAL_SKIPPED','信息',f"{filename} / {sh['name']} 第{r+1}行汇总行已排除"));continue
            d={h:decimal_value(v(h)) for h in NUMERIC}
            red=[c['cell'] for c,value in zip(cells,vals) if c.get('red_font') and value not in ('',None)]
            red_fill=[c['cell'] for c,value in zip(cells,vals) if c.get('red_fill') and value not in ('',None)]
            excluded=bool(exclude_red and red)
            item={'记录ID':f'{digest}:{sh["name"]}:{r+1}','来源':expected,'原文件':os.path.basename(filename),
                '文件SHA256':digest,'工作表':sh['name'],'原始行号':r+1,
                '来源性质':'红字人工补充（已排除）' if excluded else '去红字后保留数据（用户确认）' if exclude_red else '公司原表',
                '商品编码':as_code(v('商品')),'商品条码':as_code(v('商品条码')),'商品名称':as_code(v('商品名称')),
                '件装单位':as_code(v('件装单位')),'销售规格':as_code(v('销售规格')),
                '系统数':d['库存数量'],'初盘数':d['初盘数量'],'子转母':d['子转母商品数量'],
                '源复盘数':d['复盘数量'],'源报差异数':d['差异数量'],'成本价':d['成本价'],
                '原始差异成本':d['成本金额'],'源表备注':v('初复盘差异原因解释'),
                '原始字段':{h:v(h) for h in INPUT_HEADERS},'原始整行':vals,'单元格证据':cells,
                '原始表头':headers,'红字单元格':red,'红底单元格':red_fill,'红字排除':excluded,
                '复盘实数':None,'备注':'','单位修正量':None,'单位说明':'','处理':'保留','处理说明':'',
                '导入问题':[]}
            def add(code,level,msg):item['导入问题'].append(issue(code,level,msg,item))
            for h in SOURCE_REQUIRED:
                if h in NUMERIC:
                    if d[h] is None:add('REQUIRED_NUMBER','阻断',h+'缺失或不是有效有限数字')
                elif not as_code(v(h)):add('REQUIRED_ID','阻断',h+'缺失或格式不合法')
            for h in ('子转母商品数量','复盘数量','销售数量'):
                if v(h) not in ('',None) and d[h] is None:add('OPTIONAL_NUMBER','阻断',h+'有内容但不是有效数字')
            for h in ('商品','商品条码'):
                raw=v(h);code=as_code(raw)
                if isinstance(raw,(int,float)) and (not code or len(code)>15):add('ID_PRECISION','阻断',h+'为数值且存在小数或超过15位，须核对文本原值')
                if isinstance(raw,str) and re.fullmatch(r'[+-]?\d+(\.\d+)?[eE][+-]?\d+',raw.strip()):add('ID_SCIENTIFIC','阻断',h+'为科学计数法文本，不能猜测完整编码')
                if isinstance(raw,(int,float)):
                    fmt=cells[lookup[h]].get('number_format','')
                    if re.fullmatch('0+',fmt) and len(fmt)>len(code):add('ID_LEADING_ZERO','阻断',h+'靠格式显示前导零，须核对文本原值')
            for ce,value in zip(cells,vals):
                if ce.get('error'):add('CELL_ERROR','阻断',ce['cell']+'是Excel错误值')
                if ce.get('formula') and value is None:add('FORMULA_CACHE','阻断',ce['cell']+'公式无缓存结果，请在Excel/WPS重算副本')
            if d['初盘数量'] is not None and d['初盘数量']<0:add('NEGATIVE_COUNT','阻断','初盘实物总数为负；负差异属于正常盘亏')
            if d['成本价'] is not None and d['成本价']<0:add('NEGATIVE_PRICE','阻断','成本价为负')
            if d['库存数量'] is not None and d['库存数量']<0:add('NEGATIVE_SYSTEM','信息','系统库存为负，保留，不作为实盘数量错误')
            base=d['初盘数量']-d['库存数量'] if d['初盘数量'] is not None and d['库存数量'] is not None else None
            residual=d['差异数量']-base if d['差异数量'] is not None and base is not None else None
            item['公式残差']=residual
            item['单位待核实']=bool((d['子转母商品数量'] or 0)!=0 or (residual is not None and abs(residual)>Decimal('.0001')))
            if item['单位待核实']:add('UNIT_REVIEW','待复核','子转母非零或公式不一致；残差只供核对，不自动当作修正量')
            if all(d[h] is not None for h in ('差异数量','成本价','成本金额')):
                calc=(d['差异数量']*d['成本价']).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
                if abs(calc-d['成本金额'])>Decimal('.02'):add('AMOUNT_MISMATCH','待复核','成本金额与差异×成本价差超过0.02元')
            if red:add('RED_ANNOTATION','信息','按用户确认排除整条红字记录，原值仍留档' if excluded else '存在红字，未设置排除，请核实用途')
            if red_fill:add('RED_FILL','注意','存在红底，已保留样式；红字排除规则不自动排除只有红底的行')
            if item['源表备注'] not in ('',None):add('SOURCE_NOTE','待复核','原因栏有内容，只保留为源备注，不自动变成复盘数或审核原因')
            extras=[c+1 for c,value in enumerate(vals) if (c>=len(headers) or headers[c] not in INPUT_HEADERS) and value not in ('',None)]
            if extras:add('EXTRA_DATA','待复核',f'标准字段以外第{extras}列有内容，未参与计算，已留档')
            if item['商品条码'] and not item['商品条码'].isdigit():add('BARCODE_TEXT','注意','条码含非数字，保留原文，请与档案核对')
            dq=d['差异数量'];item['实际方向']='未知' if dq is None else '盘盈' if dq>0 else '盘亏' if dq<0 else '平账'
            if not excluded and item['实际方向'] not in ('未知','平账',expected):add('MIXED_DIRECTION','注意','与上传槽位方向不同，按差异符号分类，请核实是否同批数据')
            items.append(item)
    if not matched:raise ValueError('所有工作表均未找到必需字段：'+'、'.join(SOURCE_REQUIRED))
    return items,notes
